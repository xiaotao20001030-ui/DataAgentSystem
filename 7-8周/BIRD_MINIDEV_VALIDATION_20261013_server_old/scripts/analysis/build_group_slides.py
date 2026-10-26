#!/usr/bin/env python3
"""Build the 7-8 week group-meeting deck."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
RESULT = ROOT / "result"
FIGS = RESULT / "figures"
OUT = RESULT / "组会汇报_7-8周.pptx"

NAVY = RGBColor(0x1F, 0x2A, 0x44)
BLUE = RGBColor(0x2F, 0x5D, 0x8A)
GRAY = RGBColor(0x55, 0x55, 0x55)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x2E, 0x7D, 0x32)


def add_footer(slide, index):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12.1), Inches(0.25))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = f"7–8 周验证性实验 · 2026-10-22 · {index}"
    paragraph.font.size = Pt(9)
    paragraph.font.color.rgb = GRAY


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.65), Inches(0.35), Inches(12.0), Inches(0.6))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(26)
    paragraph.font.bold = True
    paragraph.font.color.rgb = NAVY
    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.67), Inches(0.95), Inches(12.0), Inches(0.35))
        p2 = box2.text_frame.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(12)
        p2.font.color.rgb = GRAY
    line = slide.shapes.add_shape(1, Inches(0.65), Inches(1.32), Inches(12.0), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0xD9, 0xDE, 0xE4)
    line.line.fill.background()


def add_bullets(slide, items, left, top, width, height, size=18, color=RGBColor(0x22, 0x22, 0x22)):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        paragraph.text = ("• " if level == 0 else "– ") + text
        paragraph.level = level
        paragraph.font.size = Pt(size - (2 if level else 0))
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(6)
    return box


def add_image_fit(slide, path, left, top, width, height):
    with Image.open(path) as image:
        image_width, image_height = image.size
    scale = min(width / image_width, height / image_height)
    final_width = image_width * scale
    final_height = image_height * scale
    final_left = left + (width - final_width) / 2
    final_top = top + (height - final_height) / 2
    slide.shapes.add_picture(
        str(path),
        Inches(final_left),
        Inches(final_top),
        width=Inches(final_width),
        height=Inches(final_height),
    )


def add_table(slide, rows, left, top, width, height, font_size=12):
    table_shape = slide.shapes.add_table(
        rows=len(rows),
        cols=len(rows[0]),
        left=Inches(left),
        top=Inches(top),
        width=Inches(width),
        height=Inches(height),
    )
    table = table_shape.table
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = str(value)
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(font_size)
                paragraph.font.name = "PingFang SC"
                if row_index == 0:
                    paragraph.font.bold = True
                    paragraph.font.color.rgb = WHITE
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE if row_index == 0 else (
                LIGHT if row_index % 2 == 0 else WHITE
            )
    return table


def add_note(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 title
    slide = prs.slides.add_slide(blank)
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY
    title = slide.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.0))
    paragraph = title.text_frame.paragraphs[0]
    paragraph.text = "BIRD Mini-Dev 7–8 周验证性实验"
    paragraph.font.size = Pt(38)
    paragraph.font.bold = True
    paragraph.font.color.rgb = WHITE
    subtitle = slide.shapes.add_textbox(Inches(0.95), Inches(3.15), Inches(11.0), Inches(0.6))
    p = subtitle.text_frame.paragraphs[0]
    p.text = "小模型 Schema 鲁棒性与外键检索偏差"
    p.font.size = Pt(20)
    p.font.color.rgb = RGBColor(0xD6, 0xE4, 0xF0)
    meta = slide.shapes.add_textbox(Inches(0.95), Inches(4.15), Inches(11.0), Inches(0.4))
    p = meta.text_frame.paragraphs[0]
    p.text = "500 题 · 5 模型 · 预注册门槛 + 事后重分析 + 纯列补实验 · 2026-10-22"
    p.font.size = Pt(12)
    p.font.color.rgb = RGBColor(0xC9, 0xD6, 0xE2)
    add_note(slide, "开场：这次汇报不是只说通过/不通过，而是把失败原因和下一步决策讲清楚。")

    # 2 design
    slide = prs.slides.add_slide(blank)
    add_title(slide, "预注册问题与实验矩阵", "原始假设：小模型加 20 列后，准确率下降比 DeepSeek 多 ≥10pp；且检索更易漏桥接成分")
    add_bullets(
        slide,
        [
            "500 题，11 个 SQLite 库；5 个模型：Qwen3-0.6B/1.7B/4B、Qwen2.5-Coder-1.5B、DeepSeek",
            "多给：random / synonym × 5/10/20/50；漏给：移除桥接表 103、外键列 494",
            "纯列补实验：只从 gold tables 加 20 列，不新增表（5 模型 × 1,000 条）",
            "评分：BIRD 官方 set 语义；配对统计按 question_id 聚类 bootstrap",
        ],
        0.7,
        1.65,
        6.2,
        4.6,
        size=16,
    )
    add_table(
        slide,
        [
            ["实验块", "规模"],
            ["多给扰动", "5 模型 × 9 条件 × 500"],
            ["真实近义列", "5 模型 × 427"],
            ["漏给扰动", "5 模型 × 597"],
            ["纯列 20", "5 模型 × 1,000"],
            ["检索", "500 题 × k=5/10/20/50"],
        ],
        7.2,
        1.75,
        5.3,
        3.9,
        font_size=13,
    )
    add_footer(slide, 2)
    add_note(slide, "强调：纯列补实验是为了解决原 random 条件同时加表的问题。")

    # 3 retrieval
    slide = prs.slides.add_slide(blank)
    add_title(slide, "检索：偏差在外键列，不在桥接表", "k=10 · question-only · 按问题配对 bootstrap")
    add_image_fit(slide, FIGS / "fig3_retrieval_matched.png", 0.55, 1.45, 8.0, 4.6)
    add_bullets(
        slide,
        [
            "原始指标：桥接成分并集 56.88%，比单表目标列高 28.17pp",
            "加入多表非桥接基线后：桥接表内列 -13.91pp（不显著）",
            "外键/join 列 +13.75pp（CI [+9.79, +17.73]）",
            "question+evidence：外键 +26.11pp，桥接列 +2.15pp",
        ],
        8.75,
        1.75,
        4.0,
        4.6,
        size=14,
    )
    add_footer(slide, 3)
    add_note(slide, "原 +28.17pp 主要是题目类型效应；真正稳健的是外键/join 列偏差。")

    # 4 execution
    slide = prs.slides.add_slide(blank)
    add_title(slide, "执行：原 10pp 门槛不通过", "四个小模型的 CI 上界最高 +5.2pp；可排除 +10pp")
    add_image_fit(slide, FIGS / "fig1_random_curve.png", 0.55, 1.45, 8.0, 4.5)
    add_table(
        slide,
        [
            ["模型", "下降差", "95% CI"],
            ["Qwen3-0.6B", "-1.80", "[-5.80, +2.00]"],
            ["Qwen3-1.7B", "-1.00", "[-5.40, +3.40]"],
            ["Qwen3-4B", "+0.80", "[-3.60, +5.20]"],
            ["Coder-1.5B", "-2.20", "[-7.00, +2.60]"],
        ],
        8.75,
        1.75,
        3.9,
        2.8,
        font_size=11,
    )
    add_bullets(
        slide,
        [
            "全体门槛失败，但 baseline-correct 额外失败 25.98–34.48pp",
            "说明原门槛受 baseline floor effect 影响",
        ],
        8.75,
        4.75,
        4.0,
        1.4,
        size=12,
    )
    add_footer(slide, 4)
    add_note(slide, "不要再说‘样本不够’：CI 上界已经排除 +10pp。")

    # 5 confounds
    slide = prs.slides.add_slide(blank)
    add_title(slide, "两个实现层面的混淆必须修正", "随机条件加表；5 列断崖主要来自 same_name")
    add_image_fit(slide, FIGS / "fig2_synonym_and_same_name.png", 0.55, 1.5, 8.1, 4.8)
    add_bullets(
        slide,
        [
            "add_random_20：494/500 个变体至少新增一张表，58.9% 新增列来自新表",
            "add_synonym_5：63.5% 是 same_name；same_name=0 时 Qwen3-4B 仍有 59.3%",
            "真实近义列：423/427 题引入新表，不能单独归因于词法真实性",
        ],
        8.85,
        1.8,
        3.9,
        4.2,
        size=13,
    )
    add_footer(slide, 5)
    add_note(slide, "5 列断崖不是‘近义词’本身，而是同名跨表干扰。")

    # 6 pure
    slide = prs.slides.add_slide(blank)
    add_title(slide, "纯列 20：全体门槛仍失败，机制仍显著", "只加列、不加表；5 模型 × 1,000 条")
    add_image_fit(slide, FIGS / "fig4_pure_column.png", 0.55, 1.45, 8.1, 4.7)
    add_bullets(
        slide,
        [
            "下降差：+0.60 / +4.60 / +1.80 / -3.80pp",
            "四个 CI 上界均 <10pp，全体门槛仍不通过",
            "baseline-correct：+32.35 / +34.13 / +10.00 / +20.51pp",
            "结论：原失败不是‘加表’造成的假象",
        ],
        8.85,
        1.75,
        3.9,
        4.3,
        size=13,
    )
    add_footer(slide, 6)
    add_note(slide, "纯列补实验是最小证实实验；全体指标失败，但条件机制成立。")

    # 7 decision
    slide = prs.slides.add_slide(blank)
    add_title(slide, "结论与请导师决定", "原假设未确认；失败原因已定位，下一步换指标、补样本量")
    add_table(
        slide,
        [
            ["方案", "做法", "评价"],
            ["A 50 题 pilot", "官方可见 50 题", "可立即做；0.6B 共同正确约 5 题"],
            ["B 补 85 条 SQL", "到 135 题", "最接近原计划；0.6B 仍可能 <30"],
            ["C 表级 135 题", "gold tables 全列", "研究变量变成表级"],
            ["D 收缩 RQ1", "BIRD + 可用 Spider", "风险最低，任务书兜底"],
        ],
        0.7,
        1.65,
        11.9,
        2.8,
        font_size=12,
    )
    add_bullets(
        slide,
        [
            "建议：A 先做 pilot + 向作者申请剩余 gold SQL；确认性主指标改为 baseline-correct",
            "独立决策：Qwen3-0.6B 是否退出确认性 gate？",
            "若 RQ3 要立住：补 retrieved-schema 执行条件（约 2,500 条预测）",
        ],
        0.75,
        4.75,
        11.8,
        1.6,
        size=15,
    )
    add_footer(slide, 7)
    add_note(slide, "请导师拍板三件事：Spider 样本量、0.6B 是否退出、RQ3 是否补实验。")

    # 8 closing
    slide = prs.slides.add_slide(blank)
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY
    box = slide.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.0))
    p = box.text_frame.paragraphs[0]
    p.text = "讨论"
    p.font.size = Pt(42)
    p.font.bold = True
    p.font.color.rgb = WHITE
    box = slide.shapes.add_textbox(Inches(0.95), Inches(3.7), Inches(10.5), Inches(0.8))
    p = box.text_frame.paragraphs[0]
    p.text = "原假设未确认，但机制已定位：外键检索、同名干扰、baseline 条件效应。"
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(0xD6, 0xE4, 0xF0)
    add_note(slide, "收尾：请导师决定 Spider 方案与 0.6B 去留。")

    for index, slide in enumerate(prs.slides, start=1):
        if index not in (1, 8):
            pass
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
