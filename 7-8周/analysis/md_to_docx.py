#!/usr/bin/env python3
"""Convert the project's Markdown reports to DOCX with basic table support."""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def set_east_asian_font(run, name="SimSun"):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_inline(paragraph, text, bold=False, code=False):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            set_east_asian_font(run)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9.5)
        else:
            run = paragraph.add_run(part)
            if bold:
                run.bold = True
            set_east_asian_font(run)


def is_special_line(line):
    stripped = line.strip()
    return (
        not stripped
        or stripped.startswith(("#", "|", ">", "- ", "```"))
        or stripped in {"---", "***"}
        or bool(re.match(r"^\d+\.\s+", stripped))
    )


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for col_index in range(cols):
            value = row[col_index] if col_index < len(row) else ""
            cell = table.cell(row_index, col_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, value, bold=(row_index == 0))
    doc.add_paragraph()


def main():
    args = parse_args()
    lines = args.input.read_text(encoding="utf-8").splitlines()
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    in_code = False
    paragraph_buffer = []
    index = 0

    def flush_paragraph():
        if paragraph_buffer:
            text = " ".join(item.strip() for item in paragraph_buffer)
            paragraph = doc.add_paragraph()
            add_inline(paragraph, text)
            paragraph_buffer.clear()

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            in_code = not in_code
            index += 1
            continue

        if in_code:
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(line)
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            paragraph.paragraph_format.space_after = Pt(0)
            index += 1
            continue

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        if stripped in {"---", "***"}:
            flush_paragraph()
            index += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = []
            for table_line in table_lines:
                cells = [cell.strip() for cell in table_line.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", cell) for cell in cells):
                    continue
                rows.append(cells)
            add_table(doc, rows)
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            doc.add_heading(heading.group(2).strip(), level=level)
            index += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            paragraph = doc.add_paragraph()
            add_inline(paragraph, stripped.lstrip("> ").strip())
            paragraph.paragraph_format.left_indent = Pt(18)
            index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            flush_paragraph()
            parts = [numbered.group(1).strip()]
            index += 1
            while index < len(lines) and not is_special_line(lines[index]):
                parts.append(lines[index].strip())
                index += 1
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, " ".join(parts))
            continue

        if stripped.startswith("- "):
            flush_paragraph()
            parts = [stripped[2:].strip()]
            index += 1
            while index < len(lines) and not is_special_line(lines[index]):
                parts.append(lines[index].strip())
                index += 1
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, " ".join(parts))
            continue

        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
