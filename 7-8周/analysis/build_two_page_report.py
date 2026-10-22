#!/usr/bin/env python3
"""Build the strict two-page 7-8 week report (HTML/PDF and DOCX)."""

import os
import signal
import subprocess
import time
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
RESULT = ROOT / "result"
FIGS = RESULT / "figures"
PDF = RESULT / "BIRD_MiniDev_7-8周_两页报告_v2.pdf"
DOCX = RESULT / "BIRD_MiniDev_7-8周_两页报告_v2.docx"
HTML = Path("/tmp/bird_two_page_v2.html")


def figure_uri(name):
    return (FIGS / name).resolve().as_uri()


def build_html():
    html = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8">
<style>
  @page {{ size: A4; margin: 12mm 13mm; }}
  * {{ box-sizing: border-box; }}
  body {{ font-family: "PingFang SC", "Songti SC", serif; color:#111; font-size:9.2pt; line-height:1.45; }}
  h1 {{ font-size:18pt; margin:0 0 2mm; }}
  .sub {{ color:#444; font-size:10pt; margin:0 0 1mm; }}
  .meta {{ color:#555; font-size:8pt; margin-bottom:3mm; }}
  .takeaway {{ border:1.2px solid #333; padding:3mm 4mm; margin:0 0 3mm; background:#f7f7f7; }}
  .takeaway b {{ color:#111; }}
  .takeaway ul {{ margin:1mm 0 0 4mm; padding:0; }}
  .takeaway li {{ margin:1mm 0; }}
  img {{ width:100%; display:block; margin:1mm 0 2mm; }}
  h2 {{ font-size:11.5pt; margin:2mm 0 1.5mm; border-bottom:1px solid #bbb; padding-bottom:1mm; }}
  ul {{ margin:0 0 2mm 4mm; padding:0; }}
  li {{ margin:0.8mm 0; }}
  table {{ border-collapse:collapse; width:100%; font-size:8.2pt; margin:1mm 0 2mm; }}
  th,td {{ border:1px solid #aaa; padding:1.2mm 1.6mm; vertical-align:top; }}
  th {{ background:#eee; }}
  .pagebreak {{ page-break-before: always; }}
  .footer {{ color:#666; font-size:7.2pt; margin-top:2mm; }}
</style></head><body>

<h1>BIRD Mini-Dev 7–8 周验证性实验</h1>
<div class="sub">小模型 Schema 鲁棒性与外键检索偏差 · 两页版 v2</div>
<div class="meta">数据：BIRD Mini-Dev SQLite，500 题，11 库 · 模型：Qwen3-0.6B/1.7B/4B、Qwen2.5-Coder-1.5B、DeepSeek deepseek-chat · 2026-10-22</div>

<div class="takeaway">
<b>结论先行</b>
<ul>
<li><b>原预注册 10pp 全体门槛不通过</b>：随机 20 列下降差 -2.20 到 +0.80pp，95% CI 上界最高 +5.2pp；数据可排除“小模型多下降 10pp”。</li>
<li><b>纯列 20 补实验仍不通过全体门槛</b>：下降差 +0.60/+4.60/+1.80/-3.80pp，CI 上界均 &lt;10pp；但 baseline-correct 额外失败仍为 +10.00 到 +34.13pp。</li>
<li><b>检索偏差是外键/join 列，不是桥接表</b>：控制多表题后，桥接表内列 -13.91pp，外键/join 列 +13.75pp。</li>
</ul>
</div>

<img src="{figure_uri('fig1_random_curve.png')}" alt="random curve">
<img src="{figure_uri('fig2_synonym_and_same_name.png')}" alt="synonym curve">

<div class="pagebreak"></div>

<img src="{figure_uri('fig3_retrieval_matched.png')}" alt="retrieval matched">
<img src="{figure_uri('fig4_pure_column.png')}" alt="pure column">

<h2>需要修正的三点</h2>
<ul>
<li>原 add_random_* 同时引入新表：20 列时 494/500 个变体至少新增一张表，58.9% 的新增列来自新表；它不能回答“只加无关列”的问题。</li>
<li>原 add_synonym_5 中 63.5% 是 same_name；5 列断崖主要来自同名干扰，不是纯词法近义。</li>
<li>真实近义列 423/427 题也引入新表，只有 4 题是 gold-only；不能把 41.22pp→6.32pp 全归因于“真实 vs 合成”。</li>
</ul>

<h2>请导师决定</h2>
<table>
<tr><th>方案</th><th>做法</th><th>评价</th></tr>
<tr><td>A 50 题 pilot</td><td>官方可见的 50 题，base + 20 列</td><td>可立即做，但功效不足：0.6B 共同正确约 5 题</td></tr>
<tr><td>B 补 85 条 gold SQL</td><td>到 135 题，重跑确认性实验</td><td>最接近原计划；0.6B 仍可能不足 30 题</td></tr>
<tr><td>C 表级 135 题</td><td>用 gold tables 全部列当 baseline</td><td>不需要缺失 SQL，但研究变量变成表级</td></tr>
<tr><td>D 收缩 RQ1</td><td>以 BIRD + 可用 Spider 结果做系统实证</td><td>风险最低，符合任务书兜底</td></tr>
</table>

<p><b>建议：</b>A 先做 pilot + 向作者申请剩余 gold SQL；确认性主指标改为 baseline-correct，并把 0.6B 是否退出 gate 作为独立决策。若 RQ3 要立住，需再补 retrieved-schema 执行条件（约 2,500 条预测）。</p>

<p><b>一句话：</b>原假设没有被确认，但失败原因已经定位：不是“无关列数量”，而是<b>外键检索、同名干扰和 baseline 条件效应</b>。下一步应换主指标、补样本量，而不是继续在原 10pp 门槛上加解释。</p>

<div class="footer">详细版：result/BIRD_MiniDev_Experiment_Report_v2.md / POSTHOC_REANALYSIS_2026-10-22.md；服务器结果：7-8周/BIRD_RESULTS/</div>
</body></html>"""
    HTML.write_text(html, encoding="utf-8")


def print_pdf():
    chrome = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
    if PDF.exists():
        PDF.unlink()
    process = subprocess.Popen(
        [
            chrome,
            "--headless=new",
            "--disable-gpu",
            "--no-sandbox",
            "--no-pdf-header-footer",
            "--user-data-dir=/tmp/chrome-two-page",
            "--virtual-time-budget=10000",
            f"--print-to-pdf={PDF}",
            f"file://{HTML}",
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        start_new_session=True,
    )
    deadline = time.time() + 120
    last_size = 0
    while time.time() < deadline:
        if PDF.exists():
            size = PDF.stat().st_size
            if size > 0 and size == last_size:
                break
            last_size = size
        time.sleep(1)
    try:
        os.killpg(os.getpgid(process.pid), signal.SIGTERM)
    except Exception:
        pass
    time.sleep(1)
    try:
        os.killpg(os.getpgid(process.pid), signal.SIGKILL)
    except Exception:
        pass


def set_font(run, size=9, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    run.font.size = Pt(size)
    run.bold = bold


def add_paragraph(doc, text, size=9, bold=False):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold)
    paragraph.paragraph_format.space_after = Pt(2)
    return paragraph


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.3)
    section.right_margin = Cm(1.3)

    paragraph = doc.add_paragraph()
    run = paragraph.add_run("BIRD Mini-Dev 7–8 周验证性实验")
    set_font(run, size=17, bold=True)
    paragraph.paragraph_format.space_after = Pt(2)

    add_paragraph(doc, "小模型 Schema 鲁棒性与外键检索偏差 · 两页版 v2", size=10)
    add_paragraph(
        doc,
        "数据：BIRD Mini-Dev SQLite，500 题，11 库 · 模型：Qwen3-0.6B/1.7B/4B、Qwen2.5-Coder-1.5B、DeepSeek deepseek-chat · 2026-10-22",
        size=7.5,
    )

    add_paragraph(doc, "结论先行", size=10, bold=True)
    for text in (
        "原预注册 10pp 全体门槛不通过：随机 20 列下降差 -2.20 到 +0.80pp，95% CI 上界最高 +5.2pp；数据可排除“小模型多下降 10pp”。",
        "纯列 20 补实验仍不通过全体门槛：下降差 +0.60/+4.60/+1.80/-3.80pp，CI 上界均 <10pp；但 baseline-correct 额外失败仍为 +10.00 到 +34.13pp。",
        "检索偏差是外键/join 列，不是桥接表：控制多表题后，桥接表内列 -13.91pp，外键/join 列 +13.75pp。",
    ):
        add_paragraph(doc, "• " + text, size=8.5)

    doc.add_picture(str(FIGS / "fig1_random_curve.png"), width=Cm(17.2))
    doc.add_picture(str(FIGS / "fig2_synonym_and_same_name.png"), width=Cm(17.2))

    doc.add_page_break()

    doc.add_picture(str(FIGS / "fig3_retrieval_matched.png"), width=Cm(17.2))
    doc.add_picture(str(FIGS / "fig4_pure_column.png"), width=Cm(17.2))

    add_paragraph(doc, "需要修正的三点", size=10, bold=True)
    for text in (
        "原 add_random_* 同时引入新表：20 列时 494/500 个变体至少新增一张表，58.9% 的新增列来自新表；它不能回答“只加无关列”的问题。",
        "原 add_synonym_5 中 63.5% 是 same_name；5 列断崖主要来自同名干扰，不是纯词法近义。",
        "真实近义列 423/427 题也引入新表，只有 4 题是 gold-only；不能把 41.22pp→6.32pp 全归因于“真实 vs 合成”。",
    ):
        add_paragraph(doc, "• " + text, size=8.5)

    add_paragraph(doc, "请导师决定", size=10, bold=True)
    table = doc.add_table(rows=5, cols=3)
    table.style = "Table Grid"
    rows = [
        ("方案", "做法", "评价"),
        ("A 50 题 pilot", "官方可见的 50 题，base + 20 列", "可立即做，但功效不足：0.6B 共同正确约 5 题"),
        ("B 补 85 条 gold SQL", "到 135 题，重跑确认性实验", "最接近原计划；0.6B 仍可能不足 30 题"),
        ("C 表级 135 题", "用 gold tables 全部列当 baseline", "不需要缺失 SQL，但研究变量变成表级"),
        ("D 收缩 RQ1", "以 BIRD + 可用 Spider 结果做系统实证", "风险最低，符合任务书兜底"),
    ]
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            set_font(run, size=7.5, bold=(row_index == 0))

    add_paragraph(
        doc,
        "建议：A 先做 pilot + 向作者申请剩余 gold SQL；确认性主指标改为 baseline-correct，并把 0.6B 是否退出 gate 作为独立决策。若 RQ3 要立住，需再补 retrieved-schema 执行条件（约 2,500 条预测）。",
        size=8.2,
    )
    add_paragraph(
        doc,
        "一句话：原假设没有被确认，但失败原因已经定位：不是“无关列数量”，而是外键检索、同名干扰和 baseline 条件效应。下一步应换主指标、补样本量，而不是继续在原 10pp 门槛上加解释。",
        size=8.2,
    )
    add_paragraph(
        doc,
        "详细版：result/BIRD_MiniDev_Experiment_Report_v2.md / POSTHOC_REANALYSIS_2026-10-22.md；服务器结果：7-8周/BIRD_RESULTS/",
        size=6.5,
    )
    doc.save(DOCX)


def main():
    build_html()
    print_pdf()
    build_docx()
    print(PDF)
    print(DOCX)


if __name__ == "__main__":
    main()
