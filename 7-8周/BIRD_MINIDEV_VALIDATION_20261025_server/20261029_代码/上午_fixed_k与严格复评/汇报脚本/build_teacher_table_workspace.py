#!/usr/bin/env python3
"""Build one 20-row teacher table from completed BIRD and Spider2 results.

Run with the bundled Python runtime after marking the DOCX authoring operation.
This script does not generate predictions, score SQL, invent missing results,
or render the DOCX. Render and inspect every page before delivering the file.
"""

import argparse
import hashlib
import json
import math
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


MODELS = ("qwen3_0_6b", "qwen3_1_7b", "qwen3_4b", "qwen2_5_coder_1_5b", "api_control")
MODEL_NAMES = {
    "qwen3_0_6b": "Qwen3 0.6B", "qwen3_1_7b": "Qwen3 1.7B", "qwen3_4b": "Qwen3 4B",
    "qwen2_5_coder_1_5b": "Qwen2.5 Coder 1.5B", "api_control": "API 对照",
}
HEADERS = ("实验", "模型", "n", "对照 EX %", "处理 EX %", "效应 pp", "95% CI pp")
TITLE = "BIRD 与 Spider2 五模型实验结果总表"
OUTPUT_STEM = "teacher_results_table"


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_results(path, label):
    if not path.is_file():
        raise FileNotFoundError(f"{label} results are not complete or missing: {path}")
    result = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(result, dict):
        raise ValueError(f"{label} result must be a JSON object")
    return result


def number(value, field, minimum=None, maximum=None):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"Missing or non-finite numeric result: {field}")
    if minimum is not None and value < minimum or maximum is not None and value > maximum:
        raise ValueError(f"Result outside allowed range: {field}={value}")
    return float(value)


def interval(low, high, label):
    low, high = number(low, label + " CI lower"), number(high, label + " CI upper")
    if low > high:
        raise ValueError(f"Reversed confidence interval: {label}")
    return low, high


def build_rows(bird, spider):
    """Validate the complete supplied results, then create exactly 20 rows."""
    if bird.get("question_count") != 500:
        raise ValueError("BIRD input must describe the complete 500-question cohort")
    if bird.get("fixed_ks", {}).get("130") != {"plain": 12, "fk1": 10}:
        raise ValueError("BIRD 130 must use fixed plain top-12 and FK seed top-10")
    if bird.get("scoring_semantics") == "official_mini_dev_set_equality":
        raise ValueError("This teacher table is specified for the original custom BIRD scorer, not official BIRD set scoring")
    if (spider.get("common_scored_n"), spider.get("common_gold_n"), spider.get("common_silver_n")) != (66, 15, 51):
        raise ValueError("Spider2 input must contain the completed common cohort of 66 questions, 15 gold and 51 silver")
    spider_rows = {}
    for row in spider.get("rows", []):
        if row.get("stratum") == "all" and row.get("condition") in ("random20", "synonym20"):
            key = (row["condition"], row["model"])
            if key in spider_rows:
                raise ValueError(f"Duplicate Spider2 result: {key}")
            spider_rows[key] = row
    expected_spider = {(condition, model) for condition in ("random20", "synonym20") for model in MODELS}
    if set(spider_rows) != expected_spider:
        raise ValueError("Spider2 requires all five models in both random20 and synonym20")
    rows, passing = [], {}
    for target in (130, 200):
        fixed = bird["fixed_ks"][str(target)]
        for method in ("plain", "fk1"):
            k = fixed[method]
            if isinstance(k, bool) or not isinstance(k, int) or k < 1:
                raise ValueError(f"Invalid fixed k for BIRD {target}: {method}={k}")
            summary = bird["retrieval_summary"][f"{method}_fixed_{target}"]
            number(summary["mean_schema_tokens"], f"BIRD {target} {method} mean tokens", 0)
        for model in MODELS:
            label = f"BIRD {target} {model}"
            try:
                result = bird["models"][model][str(target)]
            except KeyError as error:
                raise ValueError(f"Missing completed result: {label}") from error
            n = number(result["paired_n"], label + " n", 1, 500)
            if not n.is_integer():
                raise ValueError(f"Non-integer paired n: {label}")
            control = number(result["plain_accuracy"], label + " plain accuracy", 0, 1)
            treatment = number(result["fk1_accuracy"], label + " FK accuracy", 0, 1)
            effect = number(result["fk1_gain_pp"], label + " FK gain")
            if abs(effect - (treatment - control) * 100) > 1e-7:
                raise ValueError(f"BIRD gain disagrees with accuracies: {label}")
            low, high = interval(result["ci95_low_pp"], result["ci95_high_pp"], label)
            if result.get("plain_fixed_k", fixed["plain"]) != fixed["plain"] or result.get("fk1_fixed_k", fixed["fk1"]) != fixed["fk1"]:
                raise ValueError(f"BIRD result k differs from frozen k: {label}")
            rows.append({"experiment": f"BIRD {target}", "model": model, "n": int(n),
                         "control": control * 100, "treatment": treatment * 100, "effect": effect, "ci": (low, high)})
    for condition in ("random20", "synonym20"):
        passing[condition] = []
        api = spider_rows[(condition, "api_control")]
        api_drop = (number(api["base_accuracy"], "API baseline", 0, 1) - number(api["perturbed_accuracy"], "API perturbed", 0, 1)) * 100
        for model in MODELS:
            label = f"Spider2 {condition} {model}"
            result = spider_rows[(condition, model)]
            if result.get("n") != 66:
                raise ValueError(f"Spider2 model does not use the common 66 questions: {label}")
            control = number(result["base_accuracy"], label + " baseline", 0, 1)
            treatment = number(result["perturbed_accuracy"], label + " perturbed", 0, 1)
            effect = number(result["drop_difference_pp"], label + " extra drop")
            if abs(effect - ((control - treatment) * 100 - api_drop)) > 1e-7:
                raise ValueError(f"Spider2 extra drop disagrees with paired accuracies: {label}")
            low, high = interval(result["ci95_lower_pp"], result["ci95_upper_pp"], label)
            if model != "api_control" and effect > 10 and low > 0:
                passing[condition].append(model)
            rows.append({"experiment": f"Spider2 {condition}", "model": model, "n": 66,
                         "control": control * 100, "treatment": treatment * 100,
                         "effect": None if model == "api_control" else effect,
                         "ci": None if model == "api_control" else (low, high)})
    if len(rows) != 20:
        raise AssertionError("The consolidated table must contain exactly 20 rows")
    return rows, passing


def texts(bird, spider, passing, bird_path, spider_path, version):
    gate_passes = any(len(models) >= 2 for models in passing.values())
    gate_counts = "；".join(f"{condition} 为 {len(passing[condition])}/4" for condition in ("random20", "synonym20"))
    intro = (
        "BIRD 检验固定检索列数下外键一跳对执行准确率的影响，Spider2 检验 20% 命名扰动后小模型是否比 API 对照下降更多。"
        f"Spider2 {'达到' if gate_passes else '未达到'}预设判据（{gate_counts}）。"
    )
    caption = f"表 1 五模型合表　版本 {version}　来源 BIRD {bird_path.name} 与 Spider2 {spider_path.name}"
    notes = [
        "表中 EX 为执行准确率，pp 为百分点，n 为同题配对样本数。BIRD 的对照为普通 top-k，处理为 top-k 加外键一跳；效应 = 一跳 EX − 普通 EX，正值表示一跳更好。Spider2 的对照为原始命名，处理为对应扰动；效应 = 小模型下降 − API 下降，正值表示小模型额外下降。API 行只作参照，效应与区间记为 —。",
    ]
    token_notes = []
    for target in (130, 200):
        fixed = bird["fixed_ks"][str(target)]
        plain = bird["retrieval_summary"][f"plain_fixed_{target}"]["mean_schema_tokens"]
        fk = bird["retrieval_summary"][f"fk1_fixed_{target}"]["mean_schema_tokens"]
        token_notes.append(f"{target} 组普通 k={fixed['plain']}、一跳种子 k={fixed['fk1']}，实际平均 schema token 为 {plain:.2f}/{fk:.2f}，一跳减普通为 {fk-plain:+.2f}")
    notes.append("BIRD 长度设置：" + "；".join(token_notes) + "。130/200 是名义长度目标，k 在题间固定；列数不足 k 时取全部可用列。实际 token 不逐题相等，也不受 130/200 上限约束。200 组的固定 k 仅由 schema token 统计选定，不使用 gold 覆盖率或模型成绩。")
    notes.append("评分口径：BIRD 延用原动态预算实验的自定义 SQLite 结果比较器，浮点数四舍五入到 8 位小数；gold SQL 含 ORDER BY 时比较有序结果，否则比较保留重复次数的多重集。这不是 BIRD 官方集合相等评分。Spider2 使用官方结果比较器。两数据集的 EX 不作横向优劣排名。")
    notes.append(f"Spider2 的共同样本为 {spider['common_scored_n']} 题，其中 {spider['common_gold_n']} 题有 gold SQL，{spider['common_silver_n']} 题为 silver 结果匹配样本。小模型原始准确率较低，存在可下降空间有限的地板效应；silver 样本经 API 结果匹配筛选，存在偏向 API 的选择偏差。")
    notes.append("Spider2 预设判据要求在同一种扰动中，4 个小模型至少 2 个同时满足额外下降 >10 pp 且 95% CI 下限 >0。" + f"本次 {gate_counts}，{'通过' if gate_passes else '未通过'}该判据；未通过不构成小模型具有鲁棒性或与 API 等效的证据。")
    bird_reps = bird.get("bootstrap", {}).get("replicates")
    spider_reps = spider.get("bootstrap_resamples")
    if not isinstance(bird_reps, int) or bird_reps < 1 or not isinstance(spider_reps, int) or spider_reps < 1:
        raise ValueError("Both inputs must record their bootstrap replicate counts")
    notes.append(f"区间来自按题配对 bootstrap，BIRD {bird_reps:,} 次、Spider2 {spider_reps:,} 次。BIRD 仅纳入两方案均可评分的题，gold 评分异常按配对剔除，预测失败计错。区间未做多重比较校正。")
    source_note = f"结果来源：BIRD {bird_path.name}（SHA-256 前 12 位 {sha256(bird_path)[:12]}）；Spider2 {spider_path.name}（{sha256(spider_path)[:12]}）。"
    return intro, caption, notes, source_note


def display_cells(row):
    return [row["experiment"], MODEL_NAMES[row["model"]], str(row["n"]),
            f"{row['control']:.1f}", f"{row['treatment']:.1f}",
            "—" if row["effect"] is None else f"{row['effect']:+.1f}",
            "—" if row["ci"] is None else f"[{row['ci'][0]:+.1f}, {row['ci'][1]:+.1f}]"]


def write_markdown(path, rows, intro, caption, notes, source_note, bird_path, spider_path):
    lines = [f"# {TITLE}", "", intro, "", caption, "", "| " + " | ".join(HEADERS) + " |",
             "|---|---|---:|---:|---:|---:|---|"]
    lines.extend("| " + " | ".join(display_cells(row)) + " |" for row in rows)
    lines.extend(["", "## 表注", ""])
    for index, note in enumerate(notes, 1):
        lines.extend([f"{index}. {note}", ""])
    lines.extend([source_note, "", f"[BIRD 结果](<{bird_path.resolve()}>)；[Spider2 结果](<{spider_path.resolve()}>)。", ""])
    path.write_text("\n".join(lines), encoding="utf-8")


def font_settings(font, size, bold=False):
    font.name = "Arial"
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = RGBColor(0, 0, 0)


def style_fonts(style, size, bold=False):
    font_settings(style.font, size, bold)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for key in ("ascii", "hAnsi"):
        fonts.set(qn("w:" + key), "Arial")
    fonts.set(qn("w:eastAsia"), "Arial Unicode MS")


def shade(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    properties.append(fill)


def write_docx(path, rows, intro, caption, notes, source_note, provenance):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.7)
    styles = document.styles
    style_fonts(styles["Normal"], 10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    style_fonts(styles["Title"], 18, True)
    styles["Title"].paragraph_format.space_after = Pt(9)
    style_fonts(styles["Heading 1"], 11.5, True)
    styles["Heading 1"].paragraph_format.space_before = Pt(10)
    styles["Heading 1"].paragraph_format.space_after = Pt(5)
    style_fonts(styles["Caption"], 8.5)
    styles["Caption"].font.italic = False
    styles["Caption"].paragraph_format.space_before = Pt(5)
    styles["Caption"].paragraph_format.space_after = Pt(6)
    styles["Caption"].paragraph_format.keep_with_next = True
    document.core_properties.title = TITLE
    document.core_properties.subject = "BIRD 固定 k 外键一跳和 Spider2 命名扰动的五模型结果"
    document.core_properties.author = ""
    document.core_properties.keywords = "BIRD, Spider2, fixed k, execution accuracy"
    document.core_properties.comments = json.dumps(provenance, ensure_ascii=False)
    document.add_paragraph(TITLE, "Title")
    document.add_paragraph(intro)
    document.add_paragraph(caption, "Caption")
    table = document.add_table(rows=1, cols=len(HEADERS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (1.0, 1.5, .4, .85, .85, .9, 1.6)
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + side)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    properties.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", 80), ("bottom", 80), ("left", 65), ("right", 65)):
        element = OxmlElement("w:" + side)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    properties.append(margins)
    header = table.rows[0]
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header._tr.get_or_add_trPr().append(repeat)
    all_cells = [list(HEADERS)] + [display_cells(row) for row in rows]
    for index, values in enumerate(all_cells):
        row = header if index == 0 else table.add_row()
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for column_index, (cell, text, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index and column_index == 0:
                text = text.replace("BIRD ", "BIRD\n").replace("Spider2 ", "Spider2\n")
            if index and column_index == 1:
                text = text.replace("Qwen2.5 Coder 1.5B", "Qwen2.5 Coder\n1.5B")
            if index == 0 and column_index in (3, 4):
                text = text.replace(" EX ", "\nEX ")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            paragraph.paragraph_format.keep_with_next = False
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column_index < 2 else WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(text)
            font_settings(run.font, 9.5, index == 0)
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
            shade(cell, "E6EEF5" if index == 0 else ("FFFFFF" if index % 2 else "F7F9FB"))
    document.add_heading("表注", level=1)
    for index, note in enumerate(notes, 1):
        paragraph = document.add_paragraph(f"{index}. {note}")
        paragraph.paragraph_format.space_after = Pt(5)
        for run in paragraph.runs:
            font_settings(run.font, 9.5)
    paragraph = document.add_paragraph(source_note)
    paragraph.paragraph_format.space_before = Pt(5)
    for run in paragraph.runs:
        font_settings(run.font, 8.5)
    document.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--bird-results", type=Path, required=True)
    parser.add_argument("--spider-results", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    bird = load_results(args.bird_results, "BIRD")
    spider = load_results(args.spider_results, "Spider2")
    rows, passing = build_rows(bird, spider)
    version = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    intro, caption, notes, source_note = texts(bird, spider, passing, args.bird_results, args.spider_results, version)
    full_provenance = {"version_utc": version, "builder_sha256": sha256(Path(__file__)),
                       "bird": {"path": str(args.bird_results.resolve()), "sha256": sha256(args.bird_results)},
                       "spider2": {"path": str(args.spider_results.resolve()), "sha256": sha256(args.spider_results)}}
    (args.output_dir / f"{OUTPUT_STEM}_provenance.json").write_text(
        json.dumps(full_provenance, ensure_ascii=False, indent=2), encoding="utf-8")
    provenance = {"version_utc": version,
                  "bird": full_provenance["bird"]["sha256"][:16],
                  "spider2": full_provenance["spider2"]["sha256"][:16]}
    args.output_dir.mkdir(parents=True, exist_ok=True)
    markdown_path = args.output_dir / f"{OUTPUT_STEM}.md"
    docx_path = args.output_dir / f"{OUTPUT_STEM}.docx"
    write_markdown(markdown_path, rows, intro, caption, notes, source_note, args.bird_results, args.spider_results)
    write_docx(docx_path, rows, intro, caption, notes, source_note, provenance)
    print(json.dumps({"rows": len(rows), "markdown": str(markdown_path), "docx": str(docx_path),
                      "render_qa_required": True}, ensure_ascii=False))


if __name__ == "__main__":
    main()
